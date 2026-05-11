import random
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from functools import wraps

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.db.models import Count
from django.http import JsonResponse
from django.shortcuts import redirect, render

from .forms import AssignmentSearchForm, ExamSessionForm, FeedbackForm
from .models import Assignment, ExamSession, Faculty, Room


# ---------------------------------------------------------------------------
# Auth decorator
# ---------------------------------------------------------------------------

def admin_login_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if not request.session.get('admin_logged_in'):
            return redirect('admin_login')
        return view_func(request, *args, **kwargs)
    return _wrapped_view


# ---------------------------------------------------------------------------
# Auth views
# ---------------------------------------------------------------------------

def admin_login(request):
    if request.session.get('admin_logged_in'):
        return redirect('dashboard')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()

        # Credentials should be stored in environment variables in production.
        # See README for setup instructions.
        from django.conf import settings
        admin_user = getattr(settings, 'ADMIN_USERNAME', 'admin')
        admin_pass = getattr(settings, 'ADMIN_PASSWORD', 'admin123')

        if username == admin_user and password == admin_pass:
            request.session['admin_logged_in'] = True
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid username or password.')

    return render(request, 'core/admin_login.html')


def logout_view(request):
    request.session.pop('admin_logged_in', None)
    messages.success(request, 'You have been logged out successfully.')
    return redirect('admin_login')


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@admin_login_required
def dashboard(request):
    return render(request, 'core/dashboard.html')


# ---------------------------------------------------------------------------
# Upload & assignment generation
# ---------------------------------------------------------------------------

@admin_login_required
def upload(request):
    if request.method == 'POST':
        faculty_file = request.FILES.get('faculty_file')
        room_file = request.FILES.get('room_file')

        if not faculty_file or not room_file:
            messages.error(request, 'Both faculty and room Excel files are required.')
            return redirect('assignments')

        import openpyxl

        exam_session = None
        exam_id = request.POST.get('exam_session_id')
        if exam_id:
            try:
                exam_session = ExamSession.objects.get(id=exam_id)
            except ExamSession.DoesNotExist:
                messages.error(request, 'Selected exam session does not exist.')
                return redirect('assignments')

        # --- Faculty ---
        wb_faculty = openpyxl.load_workbook(faculty_file)
        ws_faculty = wb_faculty.active
        Faculty.objects.all().delete()

        headers = [str(cell.value) for cell in ws_faculty[1]]
        date_columns = headers[4:]  # columns after name, gender, department, email

        for row in ws_faculty.iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not row[1] or not row[2] or not row[3]:
                continue
            name, gender, department, email, *availability_marks = row
            available_dates = []
            for date_col, mark in zip(date_columns, availability_marks):
                if mark is not None and str(mark).strip().lower() == 'x':
                    if hasattr(date_col, 'strftime'):
                        available_dates.append(date_col.strftime('%Y-%m-%d'))
                    else:
                        available_dates.append(str(date_col).strip().split()[0])
            Faculty.objects.create(
                name=name,
                gender=gender,
                department=department,
                availability=','.join(available_dates),
                email=email,
            )

        # --- Rooms ---
        wb_room = openpyxl.load_workbook(room_file)
        ws_room = wb_room.active
        Room.objects.all().delete()

        for row in ws_room.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            room_number, strength, *block = row
            Room.objects.create(
                room_number=room_number,
                strength=strength or 0,
                block=block[0] if block else '',
            )

        # --- Auto-assign per date in exam session ---
        if exam_session and exam_session.end_date:
            Assignment.objects.filter(exam_session=exam_session).delete()
            rooms = list(Room.objects.all())
            faculty_list = list(Faculty.objects.all())
            current_date = exam_session.start_date

            while current_date <= exam_session.end_date:
                available_faculty = [
                    f for f in faculty_list
                    if any(
                        d.strip().split()[0] == str(current_date)
                        for d in (f.availability or '').split(',')
                        if d.strip()
                    )
                ]
                random.shuffle(available_faculty)
                for i, room in enumerate(rooms):
                    if i < len(available_faculty):
                        Assignment.objects.create(
                            date=current_date,
                            room=room,
                            invigilator=available_faculty[i],
                            role='invigilator',
                            exam_session=exam_session,
                        )
                current_date += timedelta(days=1)

        messages.success(request, 'Upload and assignment creation successful!')
        return redirect('assignments')

    return render(request, 'core/upload.html')


# ---------------------------------------------------------------------------
# Assignments view
# ---------------------------------------------------------------------------

@admin_login_required
def assignments(request):
    exam_session = None
    selected_date = None
    assignments_data = []
    remaining_faculty = []
    used_faculty_ids = set()

    # Handle reset
    if request.method == 'POST' and 'reset' in request.POST:
        Assignment.objects.all().delete()
        Faculty.objects.all().delete()
        Room.objects.all().delete()
        messages.success(request, 'All data has been reset.')
        return redirect('assignments')

    selected_date = request.GET.get('date', '').strip()

    if selected_date:
        db_assignments = Assignment.objects.filter(date=selected_date).select_related('room', 'invigilator')

        if db_assignments.exists():
            # Use already-saved assignments for this date
            assignments_data = [
                {
                    'room': a.room.room_number,
                    'invigilator': a.invigilator.name,
                    'department': a.invigilator.department,
                }
                for a in db_assignments
            ]
            used_faculty_ids = {a.invigilator.id for a in db_assignments}
        else:
            # Generate new assignments for this date (no exam session context here)
            available_faculty = list(
                Faculty.objects.annotate(num_assignments=Count('invigilations'))
                .order_by('num_assignments', 'id')
            )
            available_faculty = [
                f for f in available_faculty
                if any(
                    d.strip().split()[0] == selected_date
                    for d in (f.availability or '').split(',')
                    if d.strip()
                )
            ]
            rooms = list(Room.objects.all())
            for i, room in enumerate(rooms):
                if i < len(available_faculty):
                    faculty = available_faculty[i]
                    Assignment.objects.create(
                        date=selected_date,
                        room=room,
                        invigilator=faculty,
                        role='invigilator',
                        exam_session=None,
                    )
                    assignments_data.append({
                        'room': room.room_number,
                        'invigilator': faculty.name,
                        'department': faculty.department,
                    })
                    used_faculty_ids.add(faculty.id)

        # Remaining faculty for this date
        all_available = [
            f for f in Faculty.objects.all()
            if any(
                d.strip().split()[0] == selected_date
                for d in (f.availability or '').split(',')
                if d.strip()
            )
        ]
        remaining_faculty = [f for f in all_available if f.id not in used_faculty_ids]

    exam_sessions = ExamSession.objects.all().order_by('-start_date')

    context = {
        'assignments': assignments_data,
        'remaining_faculty': remaining_faculty,
        'faculty_list': Faculty.objects.all(),
        'has_faculty': Faculty.objects.exists(),
        'selected_date': selected_date,
        'today': str(date.today()),
        'exam_sessions': exam_sessions,
    }
    return render(request, 'core/assignments.html', context)


# ---------------------------------------------------------------------------
# History
# ---------------------------------------------------------------------------

@admin_login_required
def assignment_history(request):
    date_str = request.GET.get('date', '').strip()
    selected_date = None
    grouped_assignments = {}

    if date_str:
        try:
            selected_date = datetime.strptime(date_str, '%Y-%m-%d').date()
            qs = Assignment.objects.filter(date=selected_date).select_related('room', 'invigilator')
            # Group by room; collect all invigilators per room
            for a in qs:
                key = a.room.room_number
                if key not in grouped_assignments:
                    grouped_assignments[key] = []
                grouped_assignments[key].append(a)
        except ValueError:
            messages.error(request, 'Invalid date format.')

    return render(request, 'core/history.html', {
        'selected_date': selected_date,
        'grouped_assignments': grouped_assignments,
    })


# ---------------------------------------------------------------------------
# Manage exam sessions
# ---------------------------------------------------------------------------

@admin_login_required
def manage_exam_sessions(request):
    if request.method == 'POST':
        form = ExamSessionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Exam session added successfully.')
            return redirect('manage_exam_sessions')
    else:
        form = ExamSessionForm()

    sessions = ExamSession.objects.all().order_by('-start_date')
    return render(request, 'core/exam_sessions.html', {'form': form, 'sessions': sessions})


# ---------------------------------------------------------------------------
# Assignment search
# ---------------------------------------------------------------------------

@admin_login_required
def assignment_search(request):
    form = AssignmentSearchForm(request.GET or None)
    assignments_qs = Assignment.objects.none()

    if form.is_valid():
        exam_type = form.cleaned_data.get('exam_type')
        search_date = form.cleaned_data.get('date')
        sessions = ExamSession.objects.filter(exam_type=exam_type) if exam_type else ExamSession.objects.all()
        assignments_qs = Assignment.objects.filter(exam_session__in=sessions).select_related(
            'room', 'invigilator', 'exam_session'
        )
        if search_date:
            assignments_qs = assignments_qs.filter(date=search_date)

    return render(request, 'core/assignment_search.html', {
        'form': form,
        'assignments': assignments_qs,
    })


# ---------------------------------------------------------------------------
# Tag assignments to exam session
# ---------------------------------------------------------------------------

@admin_login_required
def tag_assignments_to_exam(request):
    untagged = Assignment.objects.filter(exam_session__isnull=True).order_by('date', 'room__room_number')
    exams = ExamSession.objects.all().order_by('-start_date')

    if request.method == 'POST':
        exam_id = request.POST.get('exam_session_id')
        selected_ids = request.POST.getlist('assignment_ids')
        if exam_id and selected_ids:
            try:
                exam = ExamSession.objects.get(id=exam_id)
                Assignment.objects.filter(id__in=selected_ids).update(exam_session=exam)
                messages.success(request, f'{len(selected_ids)} assignments tagged to {exam.name}.')
            except ExamSession.DoesNotExist:
                messages.error(request, 'Invalid exam session.')
            return redirect('tag_assignments_to_exam')

    return render(request, 'core/tag_assignments.html', {
        'assignments': untagged,
        'exam_sessions': exams,
    })


# ---------------------------------------------------------------------------
# Email notifications (manual trigger only)
# ---------------------------------------------------------------------------

def _is_valid_email(email):
    return bool(re.match(r'^[^@]+@[^@]+\.[^@]+$', email or ''))


def _notify_faculty_assignments(assignments_qs, exam_session_name):
    from django.conf import settings
    from django.core.mail import send_mail

    faculty_map = defaultdict(list)
    for a in assignments_qs:
        faculty_map[a.invigilator].append((a.room.room_number, a.date))

    sent = 0
    for faculty, duty_list in faculty_map.items():
        email = faculty.email
        if not _is_valid_email(email):
            continue
        lines = [f"Room: {room}, Date: {d.strftime('%Y-%m-%d')}" for room, d in duty_list]
        message = (
            f"Dear {faculty.name},\n\n"
            f"You have been assigned the following invigilation duties for {exam_session_name}:\n\n"
            + '\n'.join(lines)
            + '\n\nPlease check your dashboard for details.\n\nRegards,\nGLEC Invigilation System'
        )
        try:
            send_mail(
                subject=f'Invigilation Assignments — {exam_session_name}',
                message=message,
                from_email=settings.EMAIL_HOST_USER,
                recipient_list=[email],
                fail_silently=False,
            )
            sent += 1
        except Exception:
            pass  # Log to monitoring in production
    return sent


@admin_login_required
def send_assignment_notifications(request):
    if request.method != 'POST':
        messages.error(request, 'Invalid request method.')
        return redirect('assignments')

    exam_session_id = request.POST.get('exam_session_id')
    if not exam_session_id:
        messages.error(request, 'No exam session selected.')
        return redirect('assignments')

    try:
        exam_session = ExamSession.objects.get(id=exam_session_id)
    except ExamSession.DoesNotExist:
        messages.error(request, 'Invalid exam session.')
        return redirect('assignments')

    assignments_qs = Assignment.objects.filter(exam_session=exam_session).select_related('room', 'invigilator')
    if not assignments_qs.exists():
        messages.error(request, 'No assignments found for the selected exam session.')
        return redirect('assignments')

    sent = _notify_faculty_assignments(assignments_qs, exam_session.name)
    messages.success(request, f'Notification emails sent to {sent} faculty member(s) for {exam_session.name}.')
    return redirect('assignments')


# ---------------------------------------------------------------------------
# Feedback
# ---------------------------------------------------------------------------

def feedback_view(request):
    if request.method == 'POST':
        form = FeedbackForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Thank you for your feedback!')
            return redirect('feedback')
    else:
        form = FeedbackForm()
    return render(request, 'core/feedback.html', {'form': form})


# ---------------------------------------------------------------------------
# Debug / test data (safe to keep, admin-protected)
# ---------------------------------------------------------------------------

@admin_login_required
def test_data(request):
    faculty_list = Faculty.objects.all()
    rooms = Room.objects.all()

    faculty_data = [
        {
            'name': f.name,
            'department': f.department,
            'availability': f.availability,
            'availability_list': f.availability.split(',') if f.availability else [],
        }
        for f in faculty_list
    ]
    room_data = [
        {'room_number': r.room_number, 'strength': r.strength, 'block': r.block}
        for r in rooms
    ]

    return JsonResponse({
        'faculty_count': faculty_list.count(),
        'room_count': rooms.count(),
        'faculty_data': faculty_data,
        'room_data': room_data,
    })